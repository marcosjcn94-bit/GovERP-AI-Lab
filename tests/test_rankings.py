from datetime import UTC, datetime

from openpyxl import Workbook

from gov_erp.rankings import extract_rankings


def _ranking_workbook(path, population=1_000_000):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pilares+dimensoes-Detalhamento"
    sheet.cell(row=2, column=22, value="Sustentabilidade fiscal")
    sheet.cell(row=3, column=1, value="UF")
    sheet.cell(row=3, column=2, value="Municipio")
    sheet.cell(row=3, column=3, value="IBGE")
    sheet.cell(row=3, column=18, value="Populacao")
    row = [None] * 72
    row[0], row[1], row[2], row[17] = "PR", "Curitiba", 4106902, population
    row[21], row[22], row[23] = 70.0, 20, 1
    row[60], row[61], row[62] = 80.0, 5, 2
    for column, value in enumerate(row, 1):
        if value is not None:
            sheet.cell(row=4, column=column, value=value)
    workbook.save(path)
    return path


def test_import_keeps_published_score_and_pillar_values(tmp_path):
    snapshot = _ranking_workbook(tmp_path / "ranking.xlsx")
    retrieved = datetime(2026, 10, 7, tzinfo=UTC)

    records = extract_rankings(snapshot, "a" * 64, retrieved)

    assert len(records) == 1
    assert records[0]["overall_score"] == 80.0
    assert records[0]["overall_rank"] == 5
    assert records[0]["rank_change"] == 2
    assert records[0]["pillar_scores"]["Sustentabilidade fiscal"]["score"] == 70.0
    assert records[0]["source_sha256"] == "a" * 64


def test_import_rejects_municipality_outside_published_population_cutoff(tmp_path):
    snapshot = _ranking_workbook(tmp_path / "ranking.xlsx", population=79_999)

    try:
        extract_rankings(snapshot, "b" * 64, datetime(2026, 10, 7, tzinfo=UTC))
    except ValueError as error:
        assert "recorte populacional" in str(error)
    else:
        raise AssertionError("Expected CLP cutoff validation to reject the row")
