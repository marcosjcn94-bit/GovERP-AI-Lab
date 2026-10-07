from __future__ import annotations

import argparse
import hashlib
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path

import httpx
from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert

from gov_erp.database import OwnerSession
from gov_erp.models import PublicRankingSnapshot
from gov_erp.seed import FIXED_CITIES

EDITION_YEAR = 2026
SOURCE_URL = "https://clp.org.br/wp-content/uploads/2026/08/MUNICIPIOS-Planilha-do-Ranking-de-Competitividadade-dos-Municipios.xlsx"
SOURCE_PAGE = "https://conteudo.clp.org.br/relatorios-tecnicos-ranking-dos-estados-e-dos-municipios"
PILLAR_COLUMNS = (22, 25, 28, 31, 34, 37, 40, 43, 46, 49, 52, 55, 58)


def retrieve_workbook(snapshot: Path | None = None) -> tuple[Path, str, datetime]:
    snapshot_dir = Path("data/snapshots")
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    retrieved_at = datetime.now(UTC)
    if snapshot:
        path = snapshot.resolve(strict=True)
    else:
        response = httpx.get(SOURCE_URL, follow_redirects=True, timeout=60)
        response.raise_for_status()
        path = snapshot_dir / f"clp-municipal-ranking-{EDITION_YEAR}.xlsx"
        path.write_bytes(response.content)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {
        "source_url": SOURCE_URL,
        "source_page": SOURCE_PAGE,
        "retrieved_at": retrieved_at.isoformat(),
        "sha256": digest,
        "edition_year": EDITION_YEAR,
    }
    path.with_suffix(".manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return path, digest, retrieved_at


def extract_rankings(path: Path, digest: str, retrieved_at: datetime) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = next((item for item in workbook.worksheets if item.title.startswith("Pilares+")), None)
    if sheet is None:
        raise ValueError("A planilha CLP nao contem a aba de pilares e dimensoes esperada")
    headings = [sheet.cell(row=2, column=column).value for column in PILLAR_COLUMNS]
    target_codes = {city[0] for city in FIXED_CITIES}
    records: list[dict[str, object]] = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        code_value = row[2]
        if code_value is None:
            continue
        ibge_code = str(int(code_value)).zfill(7)
        if ibge_code not in target_codes:
            continue
        score, rank = row[60], row[61]
        if not isinstance(score, (float, int)) or not isinstance(rank, (float, int)):
            continue
        pillars = {}
        for name, column in zip(headings, PILLAR_COLUMNS, strict=True):
            if not name:
                continue
            value, position, delta = row[column - 1 : column + 2]
            pillars[str(name)] = {
                "score": value,
                "rank": position,
                "rank_change": delta if isinstance(delta, (int, float)) else None,
            }
        records.append(
            {
                "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"clp:{EDITION_YEAR}:{ibge_code}")),
                "edition_year": EDITION_YEAR,
                "ibge_code": ibge_code,
                "municipality_name": str(row[1]),
                "uf": str(row[0]),
                "population": int(row[17]),
                "overall_score": score,
                "overall_rank": int(rank),
                "rank_change": row[62] if isinstance(row[62], (int, float)) else None,
                "pillar_scores": pillars,
                "source_url": SOURCE_URL,
                "source_sha256": digest,
                "retrieved_at": retrieved_at,
            }
        )
    workbook.close()
    if not records:
        raise ValueError("Nenhum dos municipios-alvo foi encontrado no ranking CLP")
    if any(record["population"] < 80_000 for record in records):
        raise ValueError("A planilha inclui municipio abaixo do recorte populacional do CLP")
    return records


def import_ranking(snapshot: Path | None = None) -> dict[str, object]:
    path, digest, retrieved_at = retrieve_workbook(snapshot)
    records = extract_rankings(path, digest, retrieved_at)
    with OwnerSession.begin() as session:
        statement = insert(PublicRankingSnapshot).values(records)
        session.execute(
            statement.on_conflict_do_update(
                index_elements=["edition_year", "ibge_code"],
                set_={
                    "municipality_name": statement.excluded.municipality_name,
                    "uf": statement.excluded.uf,
                    "population": statement.excluded.population,
                    "overall_score": statement.excluded.overall_score,
                    "overall_rank": statement.excluded.overall_rank,
                    "rank_change": statement.excluded.rank_change,
                    "pillar_scores": statement.excluded.pillar_scores,
                    "source_url": statement.excluded.source_url,
                    "source_sha256": statement.excluded.source_sha256,
                    "retrieved_at": statement.excluded.retrieved_at,
                },
            )
        )
    return {
        "edition_year": EDITION_YEAR,
        "records_imported": len(records),
        "municipality_codes": [record["ibge_code"] for record in records],
        "source_sha256": digest,
        "snapshot_path": str(path),
        "source_page": SOURCE_PAGE,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--snapshot", type=Path, help="Reutiliza planilha XLSX ja baixada")
    args = parser.parse_args()
    print(json.dumps(import_ranking(args.snapshot), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
